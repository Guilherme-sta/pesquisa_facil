# app.py
from flask import Flask, request, jsonify, render_template, send_file
import requests
import os
import json
import secrets
import logging
import re
import io
import time

from bs4 import BeautifulSoup
from functools import wraps

# EXCEL
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# PDF
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet


# ==================================================
# CONFIG
# ==================================================
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = Flask(__name__, template_folder="../templates", static_folder="../static")
app.secret_key = os.urandom(24)

SUAP = "https://suap.ifpi.edu.br"
LOGIN_URL = SUAP + "/accounts/login/"
ADMIN_SERVIDORES = SUAP + "/admin/rh/servidor/"

SETOR_UO = 5

CACHE_FILE = os.path.join("data", "private", "cache_servidores_picos.json")
PROGRESS_FILE = os.path.join("data", "private", "cache_progresso_picos.json")

SESSIONS = {}
app.config["SESSIONS"] = SESSIONS


# ==================================================
# ==================================================
# AUTH
# ==================================================
def require_auth(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        token = request.headers.get("X-Session-Token") or request.cookies.get("ifpi_session")

        if not token or token not in SESSIONS:
            return jsonify({
                "error": "Sessão expirada."
            }), 401

        return f(*args, **kwargs)

    return wrapper


# ==================================================
# FRONT
# ==================================================
@app.route("/")
def home():
    return render_template("index.html")


@app.route("/doutorado")
def doutorado():
    return render_template("doutorado.html")


@app.route("/painel")
def painel():
    return render_template("index.html")


# ==================================================
# LOGIN
# ==================================================
@app.route("/api/login", methods=["POST"])
def login():
    # Credential forwarding to SUAP is disabled on public deployments by default.
    if os.getenv("ENABLE_SUAP_LOGIN", "").lower() not in {"1", "true", "yes"}:
        return jsonify({"error": "Login SUAP está desativado nesta implantação."}), 404

    dados = request.get_json(silent=True) or {}
    usuario = dados.get("usuario", "").strip()
    senha = dados.get("senha", "")
    if not usuario or not senha or len(usuario) > 200 or len(senha) > 1024:
        return jsonify({"error": "Informe usuário e senha válidos."}), 400

    try:
        sess = requests.Session()

        r = sess.get(LOGIN_URL, timeout=60)
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "html.parser")
        csrf_input = soup.find("input", {"name": "csrfmiddlewaretoken"})
        if not csrf_input or not csrf_input.get("value"):
            return jsonify({"error": "Não foi possível iniciar o acesso ao SUAP."}), 502

        payload = {
            "username": usuario,
            "password": senha,
            "csrfmiddlewaretoken": csrf_input["value"]
        }

        r = sess.post(
            LOGIN_URL,
            data=payload,
            headers={"Referer": LOGIN_URL},
            timeout=60
        )

        if r.status_code >= 400 or "logout" not in r.text.lower():
            return jsonify({
                "error": "Usuário ou senha incorretos, ou o SUAP não autorizou o acesso."
            }), 401

        token = secrets.token_hex(32)
        SESSIONS[token] = sess

        logger.info("Nova sessão SUAP iniciada.")

        response = jsonify({"success": True})
        response.set_cookie("ifpi_session", token, httponly=True, samesite="Strict", secure=request.is_secure)
        return response

    except requests.RequestException:
        logger.warning("Não foi possível conectar ao SUAP durante o login.")
        return jsonify({"error": "Não foi possível conectar ao SUAP. Tente novamente."}), 502


# ==================================================
# UTILITÁRIOS
# ==================================================
def salvar_json_seguro(arquivo, dados):
    temp = arquivo + ".tmp"

    with open(temp, "w", encoding="utf-8") as f:
        json.dump(
            dados,
            f,
            ensure_ascii=False,
            indent=2
        )

    os.replace(temp, arquivo)

def carregar_cache():
    if not os.path.exists(CACHE_FILE):
        return []
    try:
        if os.path.getsize(CACHE_FILE) == 0:
            return []

        with open(CACHE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)

    except Exception as e:
        logger.warning(f"Cache inválido: {str(e)}")
        return []

def carregar_progresso():
    if not os.path.exists(PROGRESS_FILE):
        return {"pagina": 1}

    try:
        with open(PROGRESS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)

    except:
        return {"pagina": 1}


def salvar_progresso(pagina):
    salvar_json_seguro(
        PROGRESS_FILE,
        {"pagina": pagina}
    )


def limpar_progresso():
    if os.path.exists(PROGRESS_FILE):
        os.remove(PROGRESS_FILE)


# ==================================================
# REQUISIÇÃO RESISTENTE
# ==================================================
def get_retry(sess, url, timeout=30):

    espera = 5

    while True:
        try:
            r = sess.get(url, timeout=timeout)

            if r.status_code == 200:
                return r

            logger.warning(
                f"HTTP {r.status_code}. Esperando..."
            )

        except Exception as e:
            logger.warning(
                f"Erro conexão: {str(e)}"
            )

        logger.info(
            f"Tentando novamente em {espera}s..."
        )

        time.sleep(espera)

        if espera < 60:
            espera += 5


# ==================================================
# FILTROS
# ==================================================
# ==================================================
# FILTROS
# ==================================================
# ==================================================
# FILTROS (CORRIGIDO)
# ==================================================
def extrair_ano(texto):

    texto_lower = texto.lower()

    palavras_conclusao = [
        "concluído",
        "concluido",
        "conclusão",
        "conclusao",
        "finalizado",
        "finalizada",
        "titulado",
        "defendido",
        "ano de obtenção",
        "ano de obtencao"
    ]

    if not any(p in texto_lower for p in palavras_conclusao):
        return None

    # ------------------------------------------------
    # REGRA PRINCIPAL:
    # Se existir "(Concluído):" usa o ANO ANTES do texto
    # Ex:
    # 2016 - 2024 (Concluído): Doutorado...
    # => retorna 2024
    # ------------------------------------------------
    m = re.search(
        r"(\d{4})\s*\(Conclu[ií]do\)",
        texto,
        re.I
    )

    if m:
        return int(m.group(1))

    # ------------------------------------------------
    # Também aceita intervalo finalizado
    # 2016 - 2024 (Concluído)
    # ------------------------------------------------
    m = re.search(
        r"\d{4}\s*[-–]\s*(\d{4})\s*\(Conclu[ií]do\)",
        texto,
        re.I
    )

    if m:
        return int(m.group(1))

    # ------------------------------------------------
    # Se não existir isso, usa Ano de Obtenção
    # ------------------------------------------------
    m = re.search(
        r"Ano de Obten[cç][aã]o[: ]+(\d{4})",
        texto,
        re.I
    )

    if m:
        return int(m.group(1))

    # ------------------------------------------------
    # concluído ... 2025
    # ------------------------------------------------
    m = re.search(
        r"(conclu[ií]do|conclus[aã]o|finalizado|defendido).*?(19\d{2}|20\d{2})",
        texto,
        re.I
    )

    if m:
        return int(m.group(2))

    return None


def tipo_titulo(texto):

    t = texto.lower()

    if "doutorado" in t:
        return "Doutorado"

    if "mestrado" in t:
        return "Mestrado"

    return None


def filtrar_titulos(inicio, fim):

    dados = carregar_cache()
    resultados = []

    for s in dados:

        formacao = s.get("formacao", "")

        if not formacao:
            continue

        blocos = formacao.split(" | ")

        for item in blocos:

            tipo = tipo_titulo(item)

            if not tipo:
                continue

            ano = extrair_ano(item)

            # ignora se não estiver concluído
            if ano is None:
                continue

            if inicio <= ano <= fim:

                resultados.append({
                    "nome": s["nome"],
                    "matricula": s["matricula"],
                    "email": s["email"],
                    "titulo": tipo,
                    "ano": ano,
                    "descricao": item
                })

    return resultados


# ==================================================
# BUSCA
# ==================================================
@app.route("/api/servidores")
@require_auth
def servidores():

    nome = request.args.get("nome", "").lower()

    dados = carregar_cache()

    lista = []

    for s in dados:

        texto = (
            s["nome"].lower()
            + " "
            + s["matricula"]
            + " "
            + s["email"].lower()
        )

        if nome in texto:
            lista.append(s)

    return jsonify({
        "success": True,
        "total": len(lista),
        "servidores": lista[:50]
    })


# ==================================================
# TITULOS
# ==================================================
@app.route("/api/titulos")
@require_auth
def titulos():

    inicio = int(request.args.get("inicio", 2025))
    fim = int(request.args.get("fim", 2026))

    dados = filtrar_titulos(inicio, fim)

    return jsonify({
        "success": True,
        "total": len(dados),
        "resultados": dados
    })


# ==================================================
# EXPORTAR EXCEL
# ==================================================
@app.route("/api/exportar/excel")
@require_auth
def exportar_excel():

    dados = filtrar_titulos(2025, 2026)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Titulos"

    ws.append([
        "Nome",
        "Matrícula",
        "Email",
        "Título",
        "Ano",
        "Descrição"
    ])

    for c in range(1, 7):
        cel = ws.cell(row=1, column=c)
        cel.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cel.fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )
        cel.alignment = Alignment(
            horizontal="center"
        )

    for item in dados:
        ws.append([
            item["nome"],
            item["matricula"],
            item["email"],
            item["titulo"],
            item["ano"],
            item["descricao"]
        ])

    arq = io.BytesIO()
    wb.save(arq)
    arq.seek(0)

    return send_file(
        arq,
        as_attachment=True,
        download_name="titulos.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# ==================================================
# EXPORTAR PDF
# ==================================================
@app.route("/api/exportar/pdf")
@require_auth
def exportar_pdf():

    dados = filtrar_titulos(2025, 2026)

    dados.sort(
        key=lambda x: (
            0 if x["titulo"] == "Mestrado" else 1,
            x["nome"].lower()
        )
    )

    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4)
    )

    styles = getSampleStyleSheet()

    itens = []

    itens.append(
        Paragraph(
            "Servidores titulados 2025-2026",
            styles["Title"]
        )
    )

    itens.append(Spacer(1, 12))

    tabela = [[
        "Nome",
        "Matrícula",
        "Email"
        "Título",
        "Ano"
    ]]

    for item in dados:
        tabela.append([
            item["nome"],
            item["matricula"],
            item["email"],
            item["titulo"],
            str(item["ano"])
        ])

    tb = Table(tabela)

    tb.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.darkblue),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 1, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8)
    ]))

    itens.append(tb)

    doc.build(itens)

    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="titulos.pdf",
        mimetype="application/pdf"
    )


# ==================================================
# ATUALIZAR CACHE PROFISSIONAL
# ==================================================
@app.route("/api/atualizar-cache", methods=["POST"])
@require_auth
def atualizar_cache():

    token = request.headers.get("X-Session-Token")
    sess = SESSIONS[token]

    filtro_nome = request.args.get("nome", "").strip()
    usar_lattes = request.args.get("lattes", "1") == "1"

    resultados = carregar_cache()

    matriculas_vistas = {
        x["matricula"] for x in resultados
    }

    page = carregar_progresso()["pagina"]

    try:
        while True:

            url = (
                f"{ADMIN_SERVIDORES}"
                f"?ativo__exact=1"
                f"&setor__uo={SETOR_UO}"
                f"&p={page}"
            )

            if filtro_nome:
                url += f"&q={filtro_nome}"

            logger.info(f"Lendo página {page}")

            r = get_retry(sess, url)

            soup = BeautifulSoup(
                r.text,
                "html.parser"
            )

            tabela = soup.find(
                "table",
                {"id": "result_list"}
            )

            if not tabela:
                break

            tbody = tabela.find("tbody")

            if not tbody:
                break

            linhas = tbody.find_all("tr")

            if not linhas:
                break

            novos = 0

            for tr in linhas:

                info = tr.find(
                    "td",
                    class_="field-get_info_principal"
                )

                if not info:
                    continue

                h3 = info.find("h3")

                if not h3:
                    continue

                txt = h3.get_text(
                    " ",
                    strip=True
                )

                m = re.search(
                    r"(.+?)\s*\((\d+)\)",
                    txt
                )

                if not m:
                    continue

                nome = m.group(1).strip()
                matricula = m.group(2).strip()

                if matricula in matriculas_vistas:
                    continue

                matriculas_vistas.add(matricula)
                novos += 1

                email = ""

                contato = tr.find(
                    "td",
                    class_="field-get_contatos"
                )

                if contato:
                    a = contato.find(
                        "a",
                        href=re.compile("mailto")
                    )

                    if a:
                        email = a.get_text(strip=True)

                formacao = "Sem formação"

                if usar_lattes:

                    try:
                        perfil = (
                            f"{SUAP}/rh/servidor/"
                            f"{matricula}/"
                        )

                        pr = get_retry(
                            sess,
                            perfil,
                            20
                        )

                        psoup = BeautifulSoup(
                            pr.text,
                            "html.parser"
                        )

                        links = psoup.find_all(
                            "a",
                            href=re.compile(
                                r"^/cnpq/curriculo/\d+/$"
                            )
                        )

                        href = None

                        for lk in links:

                            txtlk = lk.get_text(
                                " ",
                                strip=True
                            ).lower()

                            if (
                                "currículo" in txtlk
                                or "lattes" in txtlk
                                or lk.find("img")
                            ):
                                href = lk["href"]
                                break

                        if href:

                            urlc = SUAP + href

                            cr = get_retry(
                                sess,
                                urlc,
                                20
                            )

                            csoup = BeautifulSoup(
                                cr.text,
                                "html.parser"
                            )

                            secao = csoup.find(
                                id="formacoes-academicas"
                            )

                            if secao:

                                lista = []

                                for lf in secao.find_all("tr"):

                                    tds = lf.find_all("td")

                                    if len(tds) >= 2:

                                        periodo = tds[0].get_text(
                                            " ",
                                            strip=True
                                        )

                                        curso = tds[1].get_text(
                                            " ",
                                            strip=True
                                        )

                                        lista.append(
                                            f"{periodo}: {curso[:250]}"
                                        )

                                if lista:
                                    formacao = " | ".join(lista)

                    except Exception as e:
                        logger.warning(
                            f"Erro {nome}: {str(e)}"
                        )

                resultados.append({
                    "nome": nome,
                    "matricula": matricula,
                    "email": email,
                    "formacao": formacao
                })

                salvar_json_seguro(
                    CACHE_FILE,
                    resultados
                )

            salvar_progresso(page)

            if novos == 0:
                logger.info("Fim páginas.")
                break

            page += 1

        limpar_progresso()

        return jsonify({
            "success": True,
            "total": len(resultados)
        })

    except Exception as e:
        logger.exception("Erro cache")

        return jsonify({
            "error": str(e)
        }), 500


# ==================================================
# STATUS
# ==================================================
@app.route("/api/status")
def status():
    return jsonify({"online": True})


# ==================================================
# RUN
# ==================================================
from routes.citations import citations_bp
app.register_blueprint(citations_bp)
from routes.public_data import public_data_bp
app.register_blueprint(public_data_bp)
